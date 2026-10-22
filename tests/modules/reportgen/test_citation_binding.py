"""Citation rendering uses opaque identifiers, never a standards resolver."""

import copy
from pathlib import Path

import pytest
import yaml
from docx import Document

from assetutilities.common.reportgen.dom import DocumentDOM
from assetutilities.common.reportgen.reportgen import ReportGenerator, run


@pytest.fixture(autouse=True)
def reset_dom():
    DocumentDOM._instance = None
    yield
    DocumentDOM._instance = None


def config(tmp_path, extension="md"):
    source = Path(__file__).parents[3] / "examples/workflows/reportgen/input.yml"
    cfg = yaml.safe_load(source.read_text(encoding="utf-8"))
    cfg["important_paths"]["project_root"] = str(tmp_path)
    cfg["report_details"].update(
        target_report_type=extension, timestamped_filename=False
    )
    cfg["references"] = [
        {"code_id": "DNV-OS-E301", "revision": "example-edition", "publisher": "DNV"}
    ]
    citation = {"type": "citation", "code_id": "DNV-OS-E301", "section": "example"}
    cfg["sections_L1"] = [
        {
            "name": "Method",
            "paragraphs": [[citation]],
            "sections_L2": [
                {"name": "Detail", "paragraphs": [[copy.deepcopy(citation)]]}
            ],
        }
    ]
    return cfg


def test_binding_validation():
    from assetutilities.common.reportgen.citations import (
        parse_references,
        render_citation,
    )

    assert parse_references({}) == []
    refs = parse_references(
        {"references": [{"code_id": "DNV-OS-E301", "revision": "example"}]}
    )
    assert refs[0]["source_sibling"] == "generic"
    assert (
        render_citation({"code_id": "DNV-OS-E301", "section": "A", "page": "2"}, refs)
        == "[DNV-OS-E301 example §A p.2]"
    )
    assert render_citation({"code_id": "DNV-OS-E301"}, refs) == "[DNV-OS-E301 example]"
    with pytest.raises(ValueError, match="Unbound"):
        render_citation({"code_id": "missing"}, refs)


@pytest.mark.parametrize(
    "field,value",
    [
        ("section", ""),
        ("section", " \t"),
        ("section", None),
        ("page", ""),
        ("page", " "),
        ("page", None),
        ("page", 0),
        ("page", -1),
        ("page", "0"),
        ("page", "-2"),
        ("page", "B"),
        ("page", "NaN"),
        ("page", float("inf")),
        ("page", True),
        ("page", 1.5),
    ],
)
def test_invalid_citation_locations_fail_closed(field, value):
    from assetutilities.common.reportgen.citations import render_citation

    refs = [{"code_id": "example", "revision": "2026"}]
    with pytest.raises(ValueError, match=field):
        render_citation({"code_id": "example", field: value}, refs)


def test_citation_locations_are_trimmed():
    from assetutilities.common.reportgen.citations import render_citation

    refs = [{"code_id": "example", "revision": "2026"}]
    assert (
        render_citation(
            {"code_id": "example", "section": " A.1 ", "page": " 2 "}, refs
        )
        == "[example 2026 §A.1 p.2]"
    )


@pytest.mark.parametrize("extension", ["md", "html", "docx"])
def test_missing_publisher_is_explicit(tmp_path, extension):
    cfg = config(tmp_path, extension)
    cfg["references"][0].pop("publisher")
    run(cfg)
    output = tmp_path / "outputs" / f"sample-report.out.{extension}"
    if extension == "docx":
        text = " ".join(
            c.text for t in Document(output).tables for r in t.rows for c in r.cells
        )
    else:
        text = output.read_text(encoding="utf-8")
    assert "Not supplied" in text


def test_structured_missing_publisher_matches_legacy():
    from assetutilities.common.reportgen.citations import bind_structured_report

    report = bind_structured_report(
        {
            "references": [{"code_id": "example", "revision": "2026"}],
            "sections": [{"name": "References", "blocks": []}],
        }
    )
    assert report["sections"][0]["blocks"][0]["rows"] == [
        ["example", "Not supplied", "2026"]
    ]


def test_report_hash_ignores_output_location_but_tracks_inputs(tmp_path):
    from assetutilities.common.reportgen.citations import (
        appendix_sections,
        parse_references,
    )

    def digest_row(cfg):
        section = appendix_sections(parse_references(cfg), cfg)[-1]
        text = section.paragraphs[0]["content"]
        return next(row for row in text.splitlines() if row.startswith("| Input sha256"))

    cfg = config(tmp_path)
    original = copy.deepcopy(cfg)
    first = digest_row(cfg)
    assert cfg == original
    cfg["important_paths"] = {
        key: str(tmp_path / "relocated" / key) for key in cfg["important_paths"]
    }
    assert digest_row(cfg) == first
    cfg["sections_L1"][0]["name"] = "Changed input"
    assert digest_row(cfg) != first


@pytest.mark.parametrize(
    "references",
    [
        None,
        {},
        [None],
        [{}],
        [{"code_id": "x"}],
        [{"code_id": "x", "revision": 3}],
        [{"code_id": "x", "revision": " "}],
        [{"code_id": "x", "revision": "a"}] * 2,
    ],
)
def test_invalid_references_fail_closed(references):
    from assetutilities.common.reportgen.citations import parse_references

    with pytest.raises(ValueError):
        parse_references({"references": references})


@pytest.mark.parametrize("extension", ["md", "docx", "html"])
def test_end_to_end(tmp_path, extension):
    from assetutilities.common.reportgen.citations import DISCLOSURE_TEXT
    from assetutilities.workflow_api.envelope import input_hash

    cfg = config(tmp_path, extension)
    expected_hash = input_hash(
        {key: value for key, value in cfg.items() if key != "important_paths"}
    )
    original = copy.deepcopy(cfg)
    run(cfg)
    assert cfg == original
    output = tmp_path / "outputs" / f"sample-report.out.{extension}"
    if extension == "docx":
        doc = Document(output)
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    else:
        text = output.read_text(encoding="utf-8")
    assert text.count("[DNV-OS-E301 example-edition §example]") == 2
    for token in (
        "References",
        "Disclosure",
        "Provenance",
        DISCLOSURE_TEXT,
        expected_hash,
    ):
        assert token in text
    assert str(tmp_path) not in text
    if extension == "html":
        assert "<table>" in text
        assert "http" not in text.lower()


def test_unbound_fails_before_writing(tmp_path):
    cfg = config(tmp_path)
    cfg["references"] = []
    with pytest.raises(ValueError, match="Unbound"):
        run(cfg)
    assert not list(tmp_path.rglob("*.out.*"))


@pytest.mark.parametrize(
    "references,flag,expected",
    [
        (True, None, True),
        (True, False, False),
        (False, True, True),
        (False, None, False),
    ],
)
def test_provenance_and_repeat_initialization(tmp_path, references, flag, expected):
    cfg = config(tmp_path)
    cfg["sections_L1"] = []
    if not references:
        cfg.pop("references")
    if flag is not None:
        cfg["provenance"] = flag
    generator = ReportGenerator()
    generator.initialize_from_config(cfg)
    names = [s.name for s in generator.dom.sections]
    assert ("Provenance" in names) is expected
    generator.initialize_from_config(cfg)
    assert [s.name for s in generator.dom.sections] == names


def test_html_escapes_content_and_table(tmp_path):
    cfg = config(tmp_path, "html")
    cfg["sections_L1"] = [
        {
            "name": "<unsafe>",
            "paragraphs": [
                [
                    {"type": "text", "content": "<script>bad</script>"},
                    {
                        "type": "markdown_table",
                        "content": "| heading |\n|---|\n| <img> |",
                    },
                ]
            ],
        }
    ]
    run(cfg)
    text = (tmp_path / "outputs/sample-report.out.html").read_text(encoding="utf-8")
    assert "<script>" not in text and "<img>" not in text
    assert "&lt;script&gt;" in text and "&lt;img&gt;" in text


def test_structured_report_bindings(tmp_path):
    import json

    from assetutilities.common.reportgen.citations import DISCLOSURE_TEXT
    from assetutilities.common.reportgen.standard import write_docx

    source = (
        Path(__file__).parents[3] / "examples/workflows/reportgen/standard-report.json"
    )
    report = json.loads(source.read_text(encoding="utf-8"))
    report["references"] = [
        {"code_id": "NIST-SP-811", "revision": "2008", "publisher": "NIST"}
    ]
    report["sections"][3]["blocks"].append(
        {"type": "citation", "code_id": "NIST-SP-811", "section": "4.3, table 5"}
    )
    original = copy.deepcopy(report)
    first = write_docx(report, tmp_path / "first.docx")
    second = write_docx(report, tmp_path / "second.docx")
    assert first.read_bytes() == second.read_bytes()
    assert report == original
    document = Document(first)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "[NIST-SP-811 2008 §4.3, table 5]" in text
    assert DISCLOSURE_TEXT in text
    assert any(
        "NIST-SP-811" in c.text
        for t in document.tables
        for r in t.rows
        for c in r.cells
    )
    report["references"] = []
    with pytest.raises(ValueError, match="Unbound"):
        write_docx(report, tmp_path / "invalid.docx")
    assert not (tmp_path / "invalid.docx").exists()


@pytest.mark.parametrize("timestamped", [True, False])
def test_legacy_output_without_bindings(tmp_path, timestamped):
    cfg = config(tmp_path)
    cfg.pop("references")
    cfg["sections_L1"] = [{"name": "Summary", "paragraphs": [["Synthetic text."]]}]
    if timestamped:
        cfg["report_details"].pop("timestamped_filename")
    run(cfg)
    (output,) = (tmp_path / "outputs").glob("*.out.md")
    assert (output.name == "sample-report.out.md") is not timestamped
    text = output.read_text(encoding="utf-8")
    assert "## Summary\nSynthetic text.\n" in text
    assert "## References" not in text and "## Provenance" not in text


@pytest.mark.pdf_integration
def test_citation_survives_pdf_conversion(tmp_path):
    import json

    from PyPDF2 import PdfReader

    from assetutilities.common.reportgen.standard import convert_pdf, write_docx

    source = (
        Path(__file__).parents[3] / "examples/workflows/reportgen/standard-report.json"
    )
    report = json.loads(source.read_text(encoding="utf-8"))
    report["references"] = [
        {"code_id": "NIST-SP-811", "revision": "2008", "publisher": "NIST"}
    ]
    report["sections"][3]["blocks"].append(
        {"type": "citation", "code_id": "NIST-SP-811", "section": "4.3"}
    )
    docx = write_docx(report, tmp_path / "citation.docx")
    pdf = convert_pdf(docx, tmp_path / "citation.pdf", tmp_path / "conversion")
    text = "\n".join(page.extract_text() for page in PdfReader(pdf).pages)
    assert "NIST-SP-811 2008" in text
    assert "must be obtained from the respective publisher" in " ".join(text.split())
