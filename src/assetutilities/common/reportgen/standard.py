"""Opt-in P6 Word formatting and PDF conversion for structured reports.

The legacy report writers retain their existing input contract. This module
adds a controlled layout to the same reportgen package without interpreting
engineering results or qualifying their sources.
"""

import argparse
import io
import json
import os
import shutil
import subprocess
import tempfile
from datetime import date, datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PyPDF2 import PdfReader, PdfWriter

SECTIONS = (
    "Executive Summary",
    "Introduction",
    "Design Data",
    "Analysis Methodology",
    "Detailed Results",
    "Key Findings",
    "Limitations",
    "Recommendations",
    "References",
    "Acronyms and Symbols",
)
DESIGN_SUBSECTIONS = (
    "Site or field data",
    "Structure or equipment data",
    "Environmental data",
    "Vessel data",
    "Criteria with their source clauses",
    "Method",
    "Load cases",
)
REFERENCE_SUBSECTIONS = ("Published references", "Internal references")


def _text(value, label):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{label} requires nonempty text")
    return value


def _validate_blocks(blocks, section_number):
    if not isinstance(blocks, list) or not blocks or not isinstance(blocks[0], str):
        raise ValueError("Each section requires an opening paragraph")
    for block in blocks:
        if isinstance(block, str):
            _text(block, "Paragraph")
            continue
        if not isinstance(block, dict) or block.get("type") != "table":
            raise ValueError("Supported blocks are text and table")
        _text(block.get("caption"), "Table caption")
        columns, rows = block.get("columns"), block.get("rows")
        if not isinstance(columns, list) or not columns:
            raise ValueError("Table columns are required")
        if not isinstance(rows, list) or not rows:
            raise ValueError("Table rows are required")
        if (
            section_number not in ("3", "9", "10")
            and section_number.isdigit()
            and len(rows) > 15
        ):
            raise ValueError("Results tables above fifteen rows belong in an appendix")
        for cell in columns:
            _text(cell, "Column heading")
        for row in rows:
            if not isinstance(row, list) or len(row) != len(columns):
                raise ValueError("Table rows must match the column count")
            for cell in row:
                _text(cell, "Table cell (format values and units explicitly)")


def validate_report(report):
    """Reject incomplete layout data; scientific adequacy needs human review."""
    if not isinstance(report, dict):
        raise ValueError("Report must be an object")
    for key in ("title", "document_number", "revision"):
        _text(report.get(key), key)
    records = report.get("revisions")
    if not isinstance(records, list) or not records:
        raise ValueError("Revision records are required")
    seen = set()
    previous_date = ""
    for record in records:
        if not isinstance(record, dict):
            raise ValueError("Each revision must be an object")
        for key in ("revision", "date", "change"):
            _text(record.get(key), key)
        date.fromisoformat(record["date"])
        if record["revision"] in seen or record["date"] < previous_date:
            raise ValueError("Revision records must be unique and chronological")
        seen.add(record["revision"])
        previous_date = record["date"]
    if records[-1]["revision"] != report["revision"]:
        raise ValueError("Newest history row must match the declared revision")
    sections = report.get("sections")
    if (
        not isinstance(sections, list)
        or any(not isinstance(s, dict) for s in sections)
        or [s.get("name") for s in sections] != list(SECTIONS)
    ):
        raise ValueError("All ten P6 sections are required in standard order")
    appendices = report.get("appendices", [])
    if not isinstance(appendices, list) or len(appendices) > 26:
        raise ValueError("Appendices must be a list with at most 26 entries")
    for index, section in enumerate(sections + appendices, 1):
        number = str(index) if index <= 10 else chr(ord("A") + index - 11)
        if not isinstance(section, dict):
            raise ValueError("Each section must be an object")
        _text(section.get("name"), "Section name")
        _validate_blocks(section.get("blocks"), number)
        subs = section.get("subsections", [])
        if not isinstance(subs, list):
            raise ValueError("Subsections must be a list")
        required = (
            DESIGN_SUBSECTIONS
            if index == 3
            else REFERENCE_SUBSECTIONS
            if index == 9
            else None
        )
        if any(not isinstance(s, dict) for s in subs):
            raise ValueError("Each subsection must be an object")
        if required and [s.get("name") for s in subs] != list(required):
            raise ValueError(f"Section {number} requires the standard subsections")
        for sub in subs:
            _text(sub.get("name"), "Subsection name")
            if sub.get("subsections"):
                raise ValueError("Only two numbered heading levels are supported")
            _validate_blocks(sub.get("blocks"), number)


def _bookmark(paragraph, name, index):
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(index))
    start.set(qn("w:name"), name)
    end.set(qn("w:id"), str(index))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _contents_link(doc, text, anchor):
    paragraph = doc.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run, node = OxmlElement("w:r"), OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def _table(doc, columns, rows):
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = False
    width = (
        doc.sections[0].page_width
        - doc.sections[0].left_margin
        - doc.sections[0].right_margin
    )
    for column in table.columns:
        column.width = int(width / len(columns))
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for cells, values in [(table.rows[0].cells, columns)] + [
        (table.add_row().cells, row) for row in rows
    ]:
        for cell, value in zip(cells, values):
            cell.width = int(width / len(columns))
            cell.text = value
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    return table


def _blocks(doc, blocks, number, count):
    for block in blocks:
        if isinstance(block, str):
            doc.add_paragraph(block)
        else:
            count += 1
            _table(doc, block["columns"], block["rows"])
            doc.add_paragraph(f"Table {number}.{count} – {block['caption']}", "Caption")
    return count


def _canonical_docx(doc):
    """Fix ZIP timestamps as well as core properties for byte reproducibility."""
    raw, result = io.BytesIO(), io.BytesIO()
    doc.save(raw)
    with ZipFile(raw) as source, ZipFile(result, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            entry = ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            entry.create_system = 0
            target.writestr(entry, source.read(name))
    return result.getvalue()


def write_docx(report, destination):
    """Write a deterministic A4 report; revision history comes only from input."""
    validate_report(report)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(18)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.paragraph_format.widow_control = True
    for style in ("Heading 1", "Heading 2", "Title"):
        doc.styles[style].font.name = "Arial"
        doc.styles[style].paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.italic = False
    props = doc.core_properties
    props.title = report["title"]
    props.author = props.last_modified_by = ""
    props.revision = 1
    props.created = props.modified = datetime.combine(
        date.fromisoformat(report["revisions"][-1]["date"]), datetime.min.time()
    )
    doc.add_heading(report["title"], 0)
    doc.add_paragraph(report["document_number"])
    doc.add_paragraph("Internal")
    doc.add_heading("1 Executive Summary", 1)
    _bookmark(doc.paragraphs[-1], "section_1", 1)
    count = _blocks(doc, report["sections"][0]["blocks"], "1", 0)
    for sub_index, sub in enumerate(report["sections"][0].get("subsections", []), 1):
        doc.add_heading(f"1.{sub_index} {sub['name']}", 2)
        count = _blocks(doc, sub["blocks"], "1", count)
    doc.add_heading("Revision history", 1)
    _table(
        doc,
        ["Revision", "Date", "What changed and why"],
        [[r["revision"], r["date"], r["change"]] for r in report["revisions"]],
    )
    doc.add_heading("Contents", 1)
    all_sections = report["sections"] + report.get("appendices", [])
    for index, item in enumerate(all_sections, 1):
        label = str(index) if index <= 10 else chr(ord("A") + index - 11)
        title = (
            f"{label} {item['name']}"
            if index <= 10
            else f"Appendix {label} – {item['name']}"
        )
        _contents_link(doc, title, f"section_{label}")
    for index, item in enumerate(all_sections[1:], 2):
        number = str(index) if index <= 10 else chr(ord("A") + index - 11)
        title = (
            f"{number} {item['name']}"
            if index <= 10
            else f"Appendix {number} – {item['name']}"
        )
        doc.add_heading(title, 1)
        _bookmark(doc.paragraphs[-1], f"section_{number}", index)
        count = _blocks(doc, item["blocks"], number, 0)
        for sub_index, sub in enumerate(item.get("subsections", []), 1):
            doc.add_heading(f"{number}.{sub_index} {sub['name']}", 2)
            count = _blocks(doc, sub["blocks"], number, count)
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(_canonical_docx(doc))
    return destination


def convert_pdf(source, destination, scratch):
    """Convert trusted DOCX using installed LibreOffice and retain native output.

    PDF metadata (including XMP dates) and random trailer IDs are omitted from
    the final copy. Page content, links and outlines are imported together.
    Byte reproducibility is scoped to the same converter and font environment.
    """
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError("PDF conversion requires installed LibreOffice")
    # The shell launcher probes system temporary folders before honoring any
    # profile setting. Its installed binary can use the bootstrap socket path.
    binary = Path(executable).resolve().with_name("soffice.bin")
    if binary.is_file():
        executable = str(binary)
    source, destination, scratch = Path(source), Path(destination), Path(scratch)
    if source.suffix.lower() != ".docx" or not source.is_file():
        raise ValueError("PDF input must be an existing DOCX document")
    if source.resolve() == destination.resolve():
        raise ValueError("PDF destination must differ from the source")
    scratch.mkdir(parents=True, exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix="report-pdf-", dir=scratch)).resolve()
    directories = {
        "HOME": work / "home",
        "TMPDIR": work / "tmp",
        "XDG_CACHE_HOME": work / "cache",
        "XDG_CONFIG_HOME": work / "config",
        "XDG_DATA_HOME": work / "data",
        "XDG_STATE_HOME": work / "state",
        "XDG_RUNTIME_DIR": work / "r",
    }
    for directory in [work / "profile", *directories.values()]:
        directory.mkdir(mode=0o700)
    environment = os.environ.copy()
    environment.update({name: str(path) for name, path in directories.items()})
    environment.update(
        {
            "OMP_NUM_THREADS": "1",
        }
    )
    log_path = work / "conversion.log"
    with log_path.open("wb") as log:
        # A fresh profile requests one normal restart (LibreOffice exit 81).
        # Retain the same profile and arguments, as its launcher does, but bound
        # the retry and propagate crashes and other failures without retrying.
        for _ in range(2):
            try:
                result = subprocess.run(
                    [
                        executable,
                        f"-env:UserInstallation={(work / 'profile').as_uri()}",
                        # Relative to cwd so long caller paths fit Unix socket limits.
                        "-env:OSL_SOCKET_PATH=r",
                        "--headless",
                        "--convert-to",
                        "pdf:writer_pdf_Export",
                        "--outdir",
                        str(work.resolve()),
                        str(source.resolve()),
                    ],
                    stdout=log,
                    stderr=subprocess.STDOUT,
                    env=environment,
                    cwd=work,
                    timeout=120,
                    check=False,
                )
            except subprocess.TimeoutExpired as error:
                raise RuntimeError(
                    f"LibreOffice conversion timed out after {error.timeout} seconds; "
                    f"inspect {log_path}"
                ) from error
            log.write(f"\nConverter exit: {result.returncode}\n".encode())
            log.flush()
            if result.returncode != 81:
                break
    native = work / f"{source.stem}.pdf"
    if result.returncode or not native.is_file():
        raise RuntimeError(
            f"LibreOffice conversion failed (exit {result.returncode}); "
            f"inspect {log_path}"
        )
    reader, writer = PdfReader(native), PdfWriter()
    # Import pages and remap link/outline targets to the final page tree.
    writer.append(reader)
    writer._root_object.pop("/Metadata", None)
    writer.add_metadata({"/Producer": "assetutilities reportgen"})
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("wb") as stream:
        writer.write(stream)
    return destination


def run(config):
    """Opt-in entry point used by the existing reportgen interface."""
    details, paths = config["report_details"], config["important_paths"]
    format_name = details.get("target_report_type", "docx")
    if format_name not in ("docx", "pdf"):
        raise ValueError("P6 output supports docx and pdf")
    stub = details.get("target_file_stub", "report")
    if Path(stub).name != stub or stub in ("", ".", ".."):
        raise ValueError("Output stub must be a filename")
    if format_name == "pdf" and not config.get("scratch_path"):
        raise ValueError("PDF conversion requires scratch_path")
    output = Path(paths["project_root"]) / paths.get("outputs_path", "")
    docx = write_docx(config["report"], output / f"{stub}.out.docx")
    if format_name == "pdf":
        return convert_pdf(docx, output / f"{stub}.out.pdf", config["scratch_path"])
    return docx


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Structured report JSON")
    parser.add_argument("--output", type=Path, required=True, help="DOCX destination")
    parser.add_argument("--pdf", type=Path, help="Optional PDF destination")
    parser.add_argument(
        "--scratch", type=Path, help="Retained conversion working directory"
    )
    args = parser.parse_args()
    if args.pdf and not args.scratch:
        parser.error("--pdf requires --scratch")
    report = json.loads(args.input.read_text(encoding="utf-8"))
    write_docx(report, args.output)
    if args.pdf:
        convert_pdf(args.output, args.pdf, args.scratch)


if __name__ == "__main__":
    main()
